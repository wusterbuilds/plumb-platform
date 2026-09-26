"""CoStar comparable transactions importer."""

import csv
import io
import logging

from app.importers.base import StructuredImporter

logger = logging.getLogger(__name__)

COSTAR_COMP_HEADERS = {
    "property name", "property address", "sale price", "price/unit",
    "price/sf", "cap rate", "sale date", "buyer", "seller",
}

HEADER_MAP = {
    "property name": "property_name",
    "property address": "address",
    "address": "address",
    "sale price": "price",
    "price": "price",
    "price/unit": "price_per_unit",
    "price per unit": "price_per_unit",
    "price/sf": "price_per_sf",
    "price per sf": "price_per_sf",
    "cap rate": "cap_rate",
    "cap rate %": "cap_rate",
    "sale date": "date",
    "close date": "date",
    "closing date": "date",
    "buyer": "buyer",
    "buyer name": "buyer",
    "seller": "seller",
    "seller name": "seller",
    "property type": "property_type",
    "type": "property_type",
    "units": "units",
    "number of units": "units",
    "rentable building area": "square_footage",
    "building sf": "square_footage",
    "gla": "square_footage",
}


class CoStarCompImporter(StructuredImporter):
    data_source = "costar"
    data_type = "comps"
    supported_mimetypes = [
        "text/csv",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ]

    def detect(self, file_bytes: bytes, filename: str) -> bool:
        fname = filename.lower()
        if "costar" not in fname and "comp" not in fname:
            return False

        if fname.endswith(".csv"):
            try:
                text = file_bytes.decode("utf-8-sig", errors="replace")
                reader = csv.reader(io.StringIO(text))
                header_row = next(reader, None)
                if header_row:
                    lower_headers = {h.strip().lower() for h in header_row}
                    return len(lower_headers & COSTAR_COMP_HEADERS) >= 3
            except Exception:
                return False

        if fname.endswith((".xlsx", ".xls")):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
                ws = wb.active
                header_row = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(max_row=1))]
                wb.close()
                return len(set(header_row) & COSTAR_COMP_HEADERS) >= 3
            except Exception:
                return False

        return False

    def parse(self, file_bytes: bytes, filename: str) -> list[dict]:
        fname = filename.lower()
        if fname.endswith(".csv"):
            return self._parse_csv(file_bytes)
        elif fname.endswith((".xlsx", ".xls")):
            return self._parse_excel(file_bytes)
        return []

    def _parse_csv(self, file_bytes: bytes) -> list[dict]:
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        records = []
        for row in reader:
            mapped = {}
            for key, val in row.items():
                normalized = key.strip().lower()
                our_key = HEADER_MAP.get(normalized)
                if our_key and val and val.strip():
                    mapped[our_key] = val.strip()
            if mapped.get("address") or mapped.get("property_name"):
                records.append(mapped)
        return records

    def _parse_excel(self, file_bytes: bytes) -> list[dict]:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not rows:
            return []

        headers = [str(h or "").strip().lower() for h in rows[0]]
        records = []
        for row in rows[1:]:
            mapped = {}
            for i, val in enumerate(row):
                if i < len(headers) and val is not None:
                    our_key = HEADER_MAP.get(headers[i])
                    if our_key:
                        mapped[our_key] = str(val).strip()
            if mapped.get("address") or mapped.get("property_name"):
                records.append(mapped)
        return records

    def validate(self, records: list[dict]) -> tuple[list[dict], list[str]]:
        valid = []
        warnings = []
        for i, rec in enumerate(records):
            cap = rec.get("cap_rate", "")
            if cap:
                try:
                    cap_val = float(cap.replace("%", "").strip())
                    if cap_val > 20:
                        warnings.append(f"Row {i+1}: cap rate {cap_val}% unusually high")
                except ValueError:
                    pass

            price = rec.get("price", "")
            if price:
                try:
                    p = float(price.replace("$", "").replace(",", "").strip())
                    if p < 0:
                        warnings.append(f"Row {i+1}: negative price ${p}")
                        continue
                except ValueError:
                    pass

            if not rec.get("address") and not rec.get("property_name"):
                warnings.append(f"Row {i+1}: missing address and property name")
                continue

            valid.append(rec)
        return valid, warnings
