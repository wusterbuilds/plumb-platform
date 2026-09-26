"""Argus DCF Excel export importer."""

import io
import logging

from app.importers.base import StructuredImporter

logger = logging.getLogger(__name__)

ARGUS_SIGNALS = {"argus", "dcf", "cash flow", "noi", "exit"}


class ArgusDCFImporter(StructuredImporter):
    data_source = "argus"
    data_type = "dcf"
    supported_mimetypes = [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ]

    def detect(self, file_bytes: bytes, filename: str) -> bool:
        fname = filename.lower()
        if not fname.endswith((".xlsx", ".xls")):
            return False
        has_signal = any(s in fname for s in ARGUS_SIGNALS)
        if has_signal:
            return True

        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
            sheet_names = {s.lower() for s in wb.sheetnames}
            wb.close()
            return bool(sheet_names & {"cash flow", "dcf", "summary", "argus"})
        except Exception:
            return False

    def parse(self, file_bytes: bytes, filename: str) -> list[dict]:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        result = {"sheets": {}}

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True, max_row=200):
                rows.append([str(c) if c is not None else "" for c in row])
            result["sheets"][sheet_name] = rows

        wb.close()

        cash_flow = self._extract_cash_flow(result)
        return [cash_flow] if cash_flow else [result]

    def _extract_cash_flow(self, parsed: dict) -> dict | None:
        """Attempt to extract structured cash flow data from Argus sheets."""
        cf = {}
        for name, rows in parsed.get("sheets", {}).items():
            name_lower = name.lower()
            if any(s in name_lower for s in ("cash flow", "dcf", "summary")):
                for row in rows:
                    if not row:
                        continue
                    label = row[0].strip().lower() if row[0] else ""
                    values = [c for c in row[1:] if c.strip()]

                    if "noi" in label or "net operating" in label:
                        cf["noi_schedule"] = values
                    elif "exit" in label:
                        cf["exit_assumptions"] = values
                    elif "cap rate" in label or "terminal" in label:
                        cf["terminal_cap_rate"] = values[0] if values else ""
                    elif "discount" in label:
                        cf["discount_rate"] = values[0] if values else ""
                    elif "npv" in label or "present value" in label:
                        cf["npv"] = values[0] if values else ""
                    elif "irr" in label:
                        cf["irr"] = values[0] if values else ""

        return cf if cf else None

    def validate(self, records: list[dict]) -> tuple[list[dict], list[str]]:
        warnings = []
        if not records:
            warnings.append("No data extracted from Argus file")
        return records, warnings
